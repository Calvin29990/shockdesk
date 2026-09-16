#!/usr/bin/env python3
"""Génère les 3 templates Word de recommandation (neutres, recyclables)."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from pathlib import Path

OUT = Path(__file__).resolve().parent
YELLOW = RGBColor(0x9A, 0x67, 0x00)
NAVY = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x55, 0x55, 0x55)


def setup(doc):
    for s in doc.sections:
        s.top_margin = Cm(2.2)
        s.bottom_margin = Cm(2.0)
        s.left_margin = Cm(2.4)
        s.right_margin = Cm(2.4)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.font.color.rgb = NAVY
    pf = style.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return doc


def p(doc, text, *, bold=False, italic=False, size=12, align="left", space_after=8, color=None, space_before=0):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.space_before = Pt(space_before)
    para.paragraph_format.line_spacing = 1.15
    if align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "justify":
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    run.font.color.rgb = color or NAVY
    return para


def mixed(doc, parts, *, align="justify", space_after=10):
    """parts = list of (text, kwargs)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.line_spacing = 1.15
    if align == "justify":
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "right":
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "center":
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text, kw in parts:
        run = para.add_run(text)
        run.bold = kw.get("bold", False)
        run.italic = kw.get("italic", False)
        run.font.size = Pt(kw.get("size", 12))
        run.font.name = "Times New Roman"
        run.font.color.rgb = kw.get("color", NAVY)
    return para


def note_jaune(doc, fr=True):
    if fr:
        t = (
            "Mode d’emploi (à supprimer avant envoi) — "
            "Les passages entre [crochets] sont à compléter, corriger ou biffer. "
            "Rien n’est figé. Une page suffit. Joindre vos coordonnées (nom, titre, téléphone, e-mail)."
        )
    else:
        t = (
            "How to use (delete before sending) — "
            "Text in [brackets] is for you to complete, correct or strike out. "
            "Nothing is locked. One page is enough. Please keep your name, title, phone and email."
        )
    p(doc, t, italic=True, size=9, color=YELLOW, space_after=14)


def florence():
    doc = setup(Document())
    p(doc, "LETTRE DE RECOMMANDATION", bold=True, size=14, align="center", space_after=4)
    p(doc, "À qui de droit", italic=True, size=11, align="center", color=GREY, space_after=14)
    note_jaune(doc, True)

    p(doc, "[Ville], le [date]", align="right", space_after=16)
    mixed(
        doc,
        [
            ("Objet : ", {"bold": True}),
            ("Recommandation — Calvin MINANG, ancien stagiaire Sales", {}),
        ],
        align="left",
        space_after=14,
    )
    p(doc, "Madame, Monsieur,", space_after=10)

    p(
        doc,
        (
            "Je soussignée, Florence OUDIN, Business Partner Senior chez FinStart.io, "
            "recommande Calvin MINANG, que j’ai encadré en tant que stagiaire Sales "
            "de juin à août 2023."
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Durant ce stage, Calvin était au contact des clients et des équipes internes. "
            "Il a fait preuve de sérieux, d’autonomie et d’une vraie posture commerciale : "
            "il préparait ses échanges, suivait ses dossiers et tenait ses engagements. "
            "[Ajouter ici un fait précis si vous le souhaitez : un dossier, un client, "
            "une situation où il a été fiable.]"
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Ce que je retiens surtout, c’est sa constance. Il était ponctuel, clair dans "
            "ses comptes rendus, et agréable à manager. Il posait des questions utiles "
            "sans se dérober quand le travail était ingrat. Je n’ai eu aucun doute sur "
            "sa fiabilité. [Biffer ou reformuler librement.]"
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Calvin prépare aujourd’hui la fin de son Master à SKEMA Business School "
            "(diplôme prévu en décembre 2026) et vise un stage de fin d’études en finance "
            "à partir de janvier 2027, puis un premier poste. Cette lettre peut être "
            "jointe à une candidature de stage comme à une candidature CDI junior. "
            "Je le recommande sans réserve pour un environnement professionnel exigeant."
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Je reste disponible pour tout complément par téléphone ou par e-mail."
        ),
        align="justify",
        space_after=16,
    )
    p(doc, "Cordialement,", space_after=20)
    p(doc, "Florence OUDIN", bold=True, space_after=2)
    p(doc, "Business Partner Senior — FinStart.io", italic=True, size=11, space_after=2)
    p(doc, "Téléphone : [numéro]", size=11, space_after=0)
    p(doc, "E-mail : [adresse e-mail]", size=11, space_after=0)
    p(doc, "LinkedIn : linkedin.com/in/… [si vous le souhaitez]", size=11, space_after=0)
    doc.save(OUT / "LETTRE_RECO_FLORENCE_OUDIN.docx")


def aronne():
    doc = setup(Document())
    p(doc, "LETTER OF RECOMMENDATION", bold=True, size=14, align="center", space_after=4)
    p(doc, "To whom it may concern", italic=True, size=11, align="center", color=GREY, space_after=14)
    note_jaune(doc, False)

    p(doc, "[City], [date]", align="right", space_after=16)
    mixed(
        doc,
        [
            ("Re: ", {"bold": True}),
            ("Recommendation — Calvin MINANG, former student", {}),
        ],
        align="left",
        space_after=14,
    )
    p(doc, "Dear Sir or Madam,", space_after=10)
    p(
        doc,
        (
            "I am Alexandre ARONNE, PhD, FRM, Professor of Finance at SKEMA Business School "
            "(Belo Horizonte). I write to recommend Calvin MINANG, who was my student "
            "during his exchange semester on the Belo Horizonte campus in 2026."
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Calvin took my classes seriously. He prepared, asked precise questions, and "
            "delivered work of a quality I would expect from a student aiming at markets "
            "and risk — not from someone ticking a box. [Add the course name(s) and, if "
            "you wish, a grade or a piece of work that stood out.]"
        ),
        align="justify",
    )
    p(
        doc,
        (
            "I also saw how he works under pressure. While he was in hospital I arranged "
            "for him to sit an exam; he still showed up prepared and did the work. That "
            "episode told me more than a transcript: he does not use difficulty as an "
            "excuse. [Keep, soften or delete this paragraph — it is true, and it is yours "
            "to use.]"
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Calvin will complete his Master at SKEMA in December 2026. He is applying "
            "for a six-month end-of-studies internship in finance from January 2027, and "
            "thereafter for a first full-time role. This letter may be used for either. "
            "I recommend him without reservation."
        ),
        align="justify",
    )
    p(
        doc,
        "I am happy to provide any further information by phone or email.",
        align="justify",
        space_after=16,
    )
    p(doc, "Yours faithfully,", space_after=20)
    p(doc, "Alexandre ARONNE, PhD, FRM", bold=True, space_after=2)
    p(doc, "Professor of Finance — SKEMA Business School, Belo Horizonte", italic=True, size=11, space_after=2)
    p(doc, "Phone: [number]", size=11, space_after=0)
    p(doc, "Email: [email]", size=11, space_after=0)
    doc.save(OUT / "LETTER_RECO_ALEXANDRE_ARONNE.docx")


def delbis():
    doc = setup(Document())
    p(doc, "LETTER OF RECOMMENDATION", bold=True, size=14, align="center", space_after=4)
    p(doc, "To whom it may concern", italic=True, size=11, align="center", color=GREY, space_after=14)
    note_jaune(doc, False)

    p(doc, "[City], [date]", align="right", space_after=16)
    mixed(
        doc,
        [
            ("Re: ", {"bold": True}),
            ("Recommendation — Calvin MINANG, former student", {}),
        ],
        align="left",
        space_after=14,
    )
    p(doc, "Dear Sir or Madam,", space_after=10)
    p(
        doc,
        (
            "I am Leonardo DELBIS DE LACERDA, Professor of Finance at SKEMA Business School "
            "(Belo Horizonte) and Financial Superintendent at BDMG. I write to recommend "
            "Calvin MINANG, who was my student during his exchange semester on the Belo "
            "Horizonte campus in 2026."
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Calvin attended my courses in foreign exchange, derivatives and banking. "
            "He was engaged, precise, and clearly oriented towards how these products "
            "actually work in a market — rates, FX, the role of a bank’s book — rather "
            "than towards a purely academic reading. [Add the exact course titles and, "
            "if useful, a grade or an assignment that stood out.]"
        ),
        align="justify",
    )
    p(
        doc,
        (
            "In class he prepared, followed the Brazilian market examples, and asked "
            "questions that showed he was connecting the local product set (FX, "
            "derivatives, the banking system) to a professional project in finance. "
            "He was reliable and straightforward to teach. [Keep, edit or strike.]"
        ),
        align="justify",
    )
    p(
        doc,
        (
            "Calvin will complete his Master at SKEMA in December 2026. He is applying "
            "for a six-month end-of-studies internship in finance from January 2027, and "
            "thereafter for a first full-time role. This letter may be used for either. "
            "I recommend him without reservation."
        ),
        align="justify",
    )
    p(
        doc,
        "I am happy to provide any further information by phone or email.",
        align="justify",
        space_after=16,
    )
    p(doc, "Yours faithfully,", space_after=20)
    p(doc, "Leonardo DELBIS DE LACERDA", bold=True, space_after=2)
    p(
        doc,
        "Professor of Finance — SKEMA Business School, Belo Horizonte",
        italic=True,
        size=11,
        space_after=2,
    )
    p(doc, "Financial Superintendent — BDMG", italic=True, size=11, space_after=2)
    p(doc, "Phone: [number]", size=11, space_after=0)
    p(doc, "Email: [email]", size=11, space_after=0)
    doc.save(OUT / "LETTER_RECO_LEONARDO_DELBIS.docx")


if __name__ == "__main__":
    florence()
    aronne()
    delbis()
    print("ok")
